"""Erzeugt plan.json aus der Excel-Datei Trainingsplan_Ironman_2027.xlsx.

Aufruf: python3 tools/excel_zu_plan.py <excel-datei> <plan.json> [ab JJJJ-MM-TT]
Mit "ab" fallen alle Einheiten vor diesem Datum weg.
"""
import json, sys, datetime
import openpyxl

SPORT = {'Schwimmen': 'schwimmen', 'Velo': 'rad', 'Lauf': 'lauf', 'Kraft': 'kraft', 'Ruhe': 'ruhetag'}

def d(v):
    return v.date().isoformat() if isinstance(v, datetime.datetime) else None

def num(v):
    return round(v, 2) if isinstance(v, (int, float)) else None

ab = sys.argv[3] if len(sys.argv) > 3 else '0000-00-00'
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)

einheiten = []
for r in wb['Training'].iter_rows(min_row=5, values_only=True):
    if not isinstance(r[0], int) or not d(r[1]) or d(r[1]) < ab:
        continue
    e = {'datum': d(r[1]), 'woche': r[0], 'sport': SPORT.get(r[3], (r[3] or '').lower()), 'titel': r[4],
         'beschreibung': r[5], 'intensitaet': r[6], 'dauer': num(r[7]), 'distanz': num(r[8]), 'belastung': num(r[20])}
    einheiten.append({k: v for k, v in e.items() if v is not None})

wochen = []
for r in wb['Wochen'].iter_rows(min_row=4, values_only=True):
    if isinstance(r[0], int):
        w = {'woche': r[0], 'von': d(r[1]), 'bis': d(r[2]), 'phase': r[3], 'besonderes': r[4], 'gewichtSoll': num(r[29])}
        if w['bis'] >= ab:
            w['von'] = max(w['von'], ab)
            wochen.append({k: v for k, v in w.items() if v is not None})

einstellungen = {}
for r in wb['Anleitung'].iter_rows(values_only=True):
    if r[0] in ('Maximalpuls (HFmax)', 'Ruhepuls', 'Startwert Fitness', 'FTP Velo (Watt)'):
        einstellungen[{'Maximalpuls (HFmax)': 'hfMax', 'Ruhepuls': 'ruhepuls', 'Startwert Fitness': 'startFitness',
                       'FTP Velo (Watt)': 'ftp'}[r[0]]] = r[1]

# Wettkämpfe und Tests mit Ziel aus dem Plan; 'art' bestimmt die Prognose
WETTKAEMPFE = [
    {'datum': '2027-01-01', 'name': 'Neujahrsmarathon Zürich', 'art': 'marathon', 'ziel': '3:59–4:13 h, kein Zeitziel', 'zielMin': 245},
    {'datum': '2027-03-07', 'name': 'Halbmarathon-Test', 'art': 'halbmarathon', 'ziel': '1:40 h', 'zielMin': 100},
    {'datum': '2027-04-11', 'name': 'Zürich Marathon', 'art': 'marathon', 'ziel': '3:29 h', 'zielMin': 209},
    {'datum': '2027-06-19', 'name': 'Spiezathlon', 'art': 'olympisch', 'ziel': '1,5 km / 40 km / 10 km', 'zielMin': 165},
    {'datum': '2027-07-04', 'name': 'Ironman Switzerland', 'art': 'ironman', 'ziel': 'Schwimmen 1:20–1:35, Marathon 4:10–4:45', 'zielMin': 770},
]

plan = {
    'athlet': 'Sandro',
    'planstart': max(ab, einheiten[0]['datum']),
    'wettkaempfe': [w for w in WETTKAEMPFE if w['datum'] >= ab],
    'wettkampf': {'name': 'Ironman Switzerland', 'ort': 'Thun', 'datum': '2027-07-04'},
    'einstellungen': einstellungen,
    'wochen': wochen,
    'einheiten': einheiten,
}
with open(sys.argv[2], 'w', encoding='utf-8') as f:
    json.dump(plan, f, ensure_ascii=False, indent=1)
print(f'{len(einheiten)} Einheiten, {len(wochen)} Wochen')
